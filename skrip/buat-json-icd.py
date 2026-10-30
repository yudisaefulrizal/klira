"""Bangun referensi JSON ICD dari salinan primer SATUSEHAT; membutuhkan openpyxl."""
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

AKAR = Path(__file__).resolve().parent.parent
SUMBER = AKAR / 'dokumen/referensi primer'
TUJUAN = AKAR / 'apps/api/src/modul/klinis/data'
MANIFEST = json.loads((SUMBER / 'manifest-icd-satusehat-2010.json').read_text())
TUJUAN.mkdir(parents=True, exist_ok=True)

for sumber in MANIFEST['berkas']:
    berkas = SUMBER / sumber['nama_berkas']
    checksum = hashlib.sha256(berkas.read_bytes()).hexdigest()
    if checksum != sumber['sha256']:
        raise ValueError(f'Checksum sumber berubah: {berkas.name}')
    workbook = load_workbook(berkas, read_only=True, data_only=True)
    sheet = workbook[sumber['lembar'][0]['nama']]
    rows = sheet.iter_rows(values_only=True)
    assert tuple(next(rows)[:3]) == ('CODE', 'DISPLAY', 'VERSION')
    kode = []
    versi = next(iter(sumber['lembar'][0]['versi']))
    for row in rows:
        if row[0] is None:
            continue
        if not isinstance(row[0], str) or not isinstance(row[1], str) or row[2] != versi:
            raise ValueError(f'Baris tidak sesuai kontrak sumber: {berkas.name}')
        kode.append({'kode': row[0], 'nama': row[1]})
    workbook.close()
    assert len(kode) == sumber['lembar'][0]['jumlah_kode']
    assert len({k['kode'] for k in kode}) == len(kode)
    kode.sort(key=lambda k: k['kode'])
    output = {
        'sistem': sumber['sistem'], 'versi': versi,
        'sumber': sumber['halaman_sumber'], 'berkas_primer': sumber['nama_berkas'],
        'sha256_primer': checksum, 'jumlah_kode': len(kode), 'kode': kode,
    }
    tujuan = TUJUAN / f"{sumber['sistem']}-2010.json"
    tujuan.write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n')
    print(f'{tujuan.relative_to(AKAR)}: {len(kode)} kode')
